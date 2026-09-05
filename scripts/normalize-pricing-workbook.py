#!/usr/bin/env python3
"""Prepare corrected workbook tables from the complete original workbook (read only).

Use the artifact-tool workbook builder to author XLSX, then --import-workbook
to read that corrected XLSX into the deployable catalog. No workbook text executes.
"""
import argparse
import hashlib
import importlib.util
import json
import re
import statistics
import uuid
from collections import defaultdict
from decimal import Decimal, ROUND_CEILING
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('original', Path(__file__).with_name('extract-pricing-catalog.py'))
original = importlib.util.module_from_spec(spec)
spec.loader.exec_module(original)
NAMESPACE = uuid.UUID('bda9d7ae-9b36-59e7-a72e-9709e166f433')
CATEGORIES = {'Camera Module':'Camera', 'Camera Lens/Glass':'Camera Lens',
              'Housing/Frame':'Housing', 'Buttons/Flex':'Buttons & Flex',
              'Motherboard/IC':'Motherboard & IC', 'Taptic/Vibrator':'Vibration Motor'}
HEADERS = ['Brand','Device Family','Model','Device Variant','Parts Category','Repair Variant',
           'Parts Cost (USD)','Markup Multiplier','Labor Fee (USD)','Suggested Minimum (USD)',
           'Source Row IDs','Supplier(s)','Product URL','Device Type','Catalog ID']

def floor(cost):
    return float((Decimal(str(cost))*2+30).quantize(Decimal('.01'), rounding=ROUND_CEILING))

def split_models(value):
    # Compatibility separators inside hardware/region parentheses are not model separators.
    depth=0; start=0; result=[]
    for i,c in enumerate(value):
        if c=='(': depth+=1
        elif c==')': depth-=1
        elif c in '/,;' and depth==0:
            result.append(value[start:i].strip()); start=i+1
    result.append(value[start:].strip())
    return result

def devices(row):
    value=row['model']
    value=re.sub(r'\bw\s*/\s*(?:out|o)\s+frame\b','',value,flags=re.I)
    value=re.sub(r'\bHigh Capacity\b|\bExtended Capacity\b','',value,flags=re.I)
    if re.search(r'\bto\b|\bSeries\b',value,re.I) and row['category']=='iPhone':
        raise ValueError('Unspecified model series or compatibility range')
    brand=row['brand']; family={'iPhone':'iPhone','iPad':'iPad','Apple Watch':'Watch','Google Pixel':'Pixel'}.get(row['category'],'')
    previous=''; result=[]
    for part in split_models(value):
        part=part.strip(' -')
        explicit=re.match(r'^(Apple|Samsung|Google|Motorola|OnePlus|Xiaomi|OPPO|Oppo|Realme|Huawei|Nokia|LG)\s+',part,re.I)
        if explicit:
            brand={'oppo':'OPPO','lg':'LG'}.get(explicit[1].lower(),explicit[1].title())
            if explicit[1].lower()=='oneplus': brand='OnePlus'
            part=part[explicit.end():]; family=''
        prefix=re.match(r'^(iPhone|iPad|Apple Watch|Watch|Galaxy|Pixelbook|Pixel|Moto|Droid|Redmi|Poco|Mi|Civi|Black Shark)\s+',part,re.I)
        if prefix:
            family={'iphone':'iPhone','ipad':'iPad','apple watch':'Watch','poco':'Poco'}.get(prefix[1].lower(),prefix[1].title())
            if family in ('iPhone','iPad','Watch') and row['brand']=='Apple': brand='Apple'
            part=part[prefix.end():]
        attrs=re.findall(r'\(([^()]*)\)',part)
        model=re.sub(r'\([^()]*\)','',part).strip(' -')
        model=re.sub(r'\bpro\s*max\b','Pro Max',model,flags=re.I)
        model=re.sub(r'\bmax\b','Max',model,flags=re.I)
        model=re.sub(r'\b(\d{1,2})\s*P\b',r'\1 Pro',model) if family=='iPhone' else model
        model=re.sub(r'\s+',' ',model).strip()
        if family=='Watch':
            model=re.sub(r'^Series (SE|Ultra)\b',r'\1',model)
            model=re.sub(r'^(SE)\s*(1|2|3)$',lambda m:'SE '+{'1':'1st','2':'2nd','3':'3rd'}[m[2]]+' Gen',model)
            model=re.sub(r'^(SE \d+(?:st|nd|rd|th))$',r'\1 Gen',model)
        if model in ('4G','5G') and previous:
            model=re.sub(r'\s+[45]G$','',previous)+' '+model
        if family=='iPhone' and model in ('Pro','Pro Max','Plus','Mini'):
            gen=re.match(r'\d+',previous)
            if gen: model=gen[0]+' '+model
        if family=='Watch' and re.fullmatch(r'\d+',model): model='Series '+model
        # iPad abbreviated numeric entries refer to the standard numbered iPad.
        # Keep Air/Mini/Pro when explicitly written rather than guessing inheritance.
        if brand=='Xiaomi' and re.match(r'^K\d',model) and family=='Redmi': pass
        if brand=='Motorola' and re.match(r'^(?:G|E)\s',model) and not family: family='Moto'
        model=re.sub(r'\b(\d+)\s*mm\b',r'\1mm',model,flags=re.I)
        if not model or len(model)>100 or re.search(r'\b(?:frame|connector|assembly|glass|battery|screen|flex|capacity|cover|board|IC|all|with|without|adhesive|original|OEM|inch)\b',model,re.I):
            raise ValueError('Model contains an unresolved part description or is incomplete')
        if any(c in model for c in '/,;()') or re.search(r'\d\s*-\s*\d',model):
            raise ValueError('Unresolved compatibility shorthand or model range')
        if family=='iPhone' and not re.fullmatch(r'(?:[4-9][SCEsce]?|1\d[eE]?|X|XS|XR|Air|SE)(?:\s+(?:Pro Max|Pro|Plus|Mini|\d+(?:st|nd|rd|th) Gen))?',model):
            raise ValueError('Unrecognized iPhone model shorthand')
        if family=='Watch' and not re.match(r'^(?:Series \d+|SE|Ultra)',model):
            raise ValueError('Unrecognized Watch generation')
        # Retain generation, year, region and hardware variants. Drop known part IDs/colors.
        keep=[]
        for attr in attrs:
            if attr.lower() in original.COLORS or attr.lower() in {'jade','lemongrass'}: continue
            if re.search(r'\b(?:616-|338S|S710|PMX|ICP|mAh|QM\d|M\d{3}|GZE|GB\d|BLP|EB-|BN\d|BM\d)',attr,re.I): continue
            if re.search(r'\b(?:20\d\d|\d+mm|\d+(?:st|nd|rd|th) Gen|WiFi|Wi-Fi|Cellular|US\b|USA\b|EU\b|Global|International|Version|Verizon|XT\d|SM-|[ASGNFXTM]\d{3})',attr,re.I):
                keep.append(attr.strip())
        variants=[[]]; model_versions=[model]
        for attr in keep:
            if family=='Watch' and model=='SE' and re.fullmatch(r'\d+(?:st|nd|rd|th)(?:\s*/\s*\d+(?:st|nd|rd|th))* Gen',attr):
                model_versions=['SE '+g+' Gen' for g in re.findall(r'\d+(?:st|nd|rd|th)',attr)]
                continue
            choices=[attr]
            if re.fullmatch(r'\d+mm(?:\s*/\s*\d+mm)+',attr) or re.fullmatch(r'20\d\d(?:\s*/\s*20\d\d)+',attr):
                choices=re.split(r'\s*/\s*',attr)
            variants=[v+[choice] for v in variants for choice in choices]
        for model_version in model_versions:
            for attrs_version in variants:
                variant='; '.join(attrs_version)
                display=' '.join(x for x in (family,model_version) if x)
                if variant: display+=' ('+variant+')'
                if len(display)>255: raise ValueError('Device label exceeds supported length')
                result.append({'brand':brand,'family':family,'modelNumber':model_version,'variant':variant,'model':display,
                               'modelId':str(uuid.uuid5(NAMESPACE,brand+'|'+display.casefold())),
                               'deviceType':original.device_type(row['category'],display)})
        previous=model
    return list({(r['brand'],r['model']):r for r in result}.values())

def prepare(source, output):
    catalog=original.extract(source)
    groups=defaultdict(list); review=[]; accepted=set(); canonical_devices={}
    for row in catalog['rows']:
        reasons=list(row['exclusionReasons'])
        if re.search(r'\b(?:tempered glass|screen protector|camera protector|adhesive|sticker|bracket only|retaining bracket)\b',row['sourceProduct'],re.I):
            reasons.append('Accessory or mounting material rather than a complete replacement part')
        parsed=[]
        if not reasons:
            try: parsed=devices(row)
            except ValueError as error: reasons.append(str(error))
        if len(parsed)==1 and parsed[0]['family']=='iPhone':
            product_models=re.findall(r'\biPhone\s+((?:1\d[eE]?|[4-9][sScC]?|XS|XR|X|Air)(?:\s+(?:Pro Max|Pro|Plus|Mini))?)\b',row['sourceProduct'],re.I)
            if len(product_models)==1 and product_models[0].casefold()!=parsed[0]['modelNumber'].casefold():
                reasons.append('Source model conflicts with the model named in the product title')
        if reasons:
            review.append([row['sourceId'],row['brand'],row['originalModel'],row['originalRepairType'],row['partsCost'],'; '.join(reasons),row['sourceUrl']])
            continue
        accepted.add(row['sourceId'])
        effective_category='Camera Lens/Glass' if row['repairType']=='Camera Lens/Glass' else row['originalRepairType']
        category=CATEGORIES.get(effective_category,effective_category)
        for device in parsed:
            device=canonical_devices.setdefault(device['modelId'],device)
            key=(device['brand'],device['family'],device['modelNumber'],device['variant'],category,row.get('repairVariant') or '')
            groups[key].append((row,device))
    prices=[]
    for i,(key,entries) in enumerate(sorted(groups.items()),1):
        rows=[x[0] for x in entries]; device=entries[0][1]
        cost=float(Decimal(str(statistics.median(r['partsCost'] for r in rows))).quantize(Decimal('.0001')))
        prices.append([*key,cost,2,30,floor(cost),','.join(str(r['sourceId']) for r in rows),
                       '; '.join(sorted({s for r in rows for s in r['suppliers']})),
                       rows[0]['sourceUrl'],device['deviceType'],i])
    data={'sourceWorkbook':source.name,'sourceSha256':hashlib.sha256(source.read_bytes()).hexdigest(),
          'headers':HEADERS,'prices':prices,'review':review,
          'sources':[[r['sourceId'],r['category'],r['originalModel'],r['originalRepairType'],r['partsCost'],
                      '; '.join(r['suppliers']),r['sourceProduct'],r['sourceUrl'],r['lowPrice'],r['highPrice'],r['listingsCompared']] for r in catalog['rows']],
          'summary':{'sourceRows':len(catalog['rows']),'acceptedSourceRows':len(accepted),'reviewRows':len(review),
                     'prices':len(prices),'models':len({tuple(p[:4]) for p in prices})}}
    assert accepted|{r[0] for r in review}=={r['sourceId'] for r in catalog['rows']}
    output.write_text(json.dumps(data,ensure_ascii=False,indent=2))
    print(json.dumps(data['summary']))

def import_workbook(path, output):
    wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
    sheet=wb['Parts Cost Database']; rows=[]
    assert list(next(sheet.iter_rows(min_row=4,max_row=4,values_only=True)))==HEADERS
    for values in sheet.iter_rows(min_row=5,values_only=True):
        if values[0] is None: continue
        brand,family,model,variant,category,repair,cost,markup,labor,minimum,sourceids,suppliers,url,device_type,id=values
        family=family or ''; variant=variant or ''; repair=repair or ''
        assert markup==2 and labor==30 and abs(minimum-floor(cost))<1e-7
        assert not any(x in str(model) for x in '/,;()')
        display=' '.join(x for x in (family,str(model)) if x)
        if variant: display+=' ('+variant+')'
        issue=category
        if category=='Screen' and repair: issue=repair+' Screen'
        if category=='Camera': issue=(repair+' Camera') if repair else 'Camera'
        rows.append({'sourceId':id,'sourceSheetRow':id+4,'category':family or brand,'brand':brand,
                     'deviceType':device_type,'model':display,'modelNumber':str(model),'deviceFamily':family,
                     'deviceVariant':variant,'modelId':str(uuid.uuid5(NAMESPACE,brand+'|'+display.casefold())),
                     'repairType':issue,'partsCategory':category,'repairVariant':repair or None,
                     'difficulty':'Fixed','partsCost':cost,'markupMultiplier':markup,'laborFee':labor,
                     'sourceSuggestedPrice':minimum,'catalogEligible':True,'originalSourceRowIds':[int(x) for x in str(sourceids).split(',')],
                     'suppliers':str(suppliers).split('; '),'sourceUrl':url})
    sources=list(wb['Source Listings'].iter_rows(min_row=5,values_only=True))
    result={'version':2,'sourceWorkbook':path.name,'sourceSha256':hashlib.sha256(path.read_bytes()).hexdigest(),
            'currency':'USD','formula':'ROUNDUP(partsCost * 2 + 30, 2)','laborFees':{'Fixed':30},
            'aggregationPolicy':'Median of eligible source row medians per individual device and repair variant',
            'originalSourceCount':len(sources),'reviewSourceCount':sum(1 for r in wb['Needs Review'].iter_rows(min_row=5,values_only=True) if r[0] is not None),
            'rows':rows}
    output.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n')
    print(json.dumps({'importedWorkbookRows':len(rows),'models':len({r['modelId'] for r in rows})}))

if __name__=='__main__':
    parser=argparse.ArgumentParser(); parser.add_argument('input',type=Path); parser.add_argument('output',type=Path)
    parser.add_argument('--import-workbook',action='store_true'); args=parser.parse_args()
    (import_workbook if args.import_workbook else prepare)(args.input,args.output)
